#!/usr/bin/env python3
"""Moyenne de l'UE calculée par Concret, pour les indicateurs notés dont la source ne publie pas de moyenne européenne.

Règle (la même sur toutes les pages) : la France est toujours comparée à la moyenne de l'UE. Si la source la publie
(Eurostat), on prend celle-là. Sinon, Concret la calcule ici, avec la même source et la même définition que pour la
France : total des pays de l'UE (27 pays actuels) divisé par leur total (population, places, PIB), ou, pour un taux
déjà en %, moyenne des pays pondérée par leur population. Aucun pays voisin seul ne sert de référence.

Pays retenus : la moyenne d'une année porte sur les pays de l'UE qui publient le chiffre cette année-là ; son évolution
(ce qui est noté) porte sur les pays qui le publient l'année et la précédente, pour qu'un pays qui entre ou sort ne la fasse
pas bouger. Ces pays doivent faire au moins 70 % de la population de l'UE (SEUIL) ; sinon l'année n'a pas de moyenne et
n'est pas comptée (la page dit pourquoi). On ne relie pas deux tables différentes de la source (périodes).

  python3 outils/ue_calc.py              calcule tout et compare à data.json, sans rien écrire
  python3 outils/ue_calc.py --ecrire     écrit ue_calc (moyenne), ue_calc_ev (évolution), ue_calc_info et ue_calc_note
  python3 outils/ue_calc.py --seul securite.rob,sante.alc   seulement ces indicateurs

Code de sortie 1 si une source est injoignable : les indicateurs de cette source gardent leurs valeurs (rien n'est écrit
pour eux), les autres sont écrits, et la prochaine vérification n'avance pas."""
import csv, io, json, os, re, subprocess, sys, time
from concurrent.futures import ThreadPoolExecutor
from datetime import date

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
SEUIL = 70.0
UE = ['BE', 'BG', 'CZ', 'DK', 'DE', 'EE', 'IE', 'EL', 'ES', 'FR', 'HR', 'IT', 'CY', 'LV', 'LT', 'LU', 'HU', 'MT', 'NL', 'AT', 'PL', 'PT', 'RO', 'SI', 'SK', 'FI', 'SE']
ISO3 = dict(zip(UE, ['BEL', 'BGR', 'CZE', 'DNK', 'DEU', 'EST', 'IRL', 'GRC', 'ESP', 'FRA', 'HRV', 'ITA', 'CYP', 'LVA', 'LTU', 'LUX', 'HUN', 'MLT', 'NLD', 'AUT', 'POL', 'PRT', 'ROU', 'SVN', 'SVK', 'FIN', 'SWE']))
NOM_EN = dict(zip(UE, ['Belgium', 'Bulgaria', 'Czechia', 'Denmark', 'Germany', 'Estonia', 'Ireland', 'Greece', 'Spain', 'France', 'Croatia', 'Italy', 'Cyprus', 'Latvia', 'Lithuania', 'Luxembourg', 'Hungary', 'Malta', 'Netherlands', 'Austria', 'Poland', 'Portugal', 'Romania', 'Slovenia', 'Slovakia', 'Finland', 'Sweden']))
NOM_FR = dict(zip(UE, ['Belgique', 'Bulgarie', 'Tchéquie', 'Danemark', 'Allemagne', 'Estonie', 'Irlande', 'Grèce', 'Espagne', 'France', 'Croatie', 'Italie', 'Chypre', 'Lettonie', 'Lituanie', 'Luxembourg', 'Hongrie', 'Malte', 'Pays-Bas', 'Autriche', 'Pologne', 'Portugal', 'Roumanie', 'Slovénie', 'Slovaquie', 'Finlande', 'Suède']))
EST = 'https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/'


def get(url, essais=4):
    # curl plutôt qu'urllib : il suit le proxy et les certificats de l'environnement de la tâche planifiée.
    for i in range(essais):
        r = subprocess.run(['curl', '-sSfL', '--max-time', '240', '-A', 'Mozilla/5.0 (concret2027.fr)', url], capture_output=True)
        if not r.returncode:
            return r.stdout
        time.sleep(2 ** (i + 1))
    raise IOError('téléchargement impossible (%s) : %s' % (url, r.stderr.decode('utf-8', 'replace').strip()))


def jour(d):
    return '%d %s %d' % (d.day, ['janvier', 'février', 'mars', 'avril', 'mai', 'juin', 'juillet', 'août', 'septembre', 'octobre', 'novembre', 'décembre'][d.month - 1], d.year)


def eurostat(table, filtres):
    """{pays: {année: valeur}} d'une table Eurostat (une seule valeur par pays et par année une fois filtrée)."""
    d = json.loads(get(EST + table + '?format=JSON&' + filtres))
    dims, tailles = d['id'], d['size']
    inv = {k: {i: c for c, i in d['dimension'][k]['category']['index'].items()} for k in dims}
    out = {}
    for idx, v in d['value'].items():
        idx, cle = int(idx), {}
        for k, s in zip(reversed(dims), reversed(tailles)):
            cle[k] = inv[k][idx % s]; idx //= s
        if cle['geo'] in UE and v is not None:
            out.setdefault(cle['geo'], {})[int(cle['time'])] = v
    return out


_pop = {}
def population():
    """Population au 1er janvier (Eurostat demo_gind), qui sert de poids et de mesure de la couverture."""
    if not _pop:
        _pop.update(eurostat('demo_gind', 'indic_de=JAN&sinceTimePeriod=1993'))
    return _pop


def pop_banque_mondiale():
    d = json.loads(get('https://api.worldbank.org/v2/country/%s/indicator/SP.POP.TOTL?format=json&date=2000:2030&per_page=2000' % ';'.join(ISO3[g] for g in UE)))
    inv = {v: k for k, v in ISO3.items()}
    out = {}
    for x in d[1]:
        if x['value'] and x['countryiso3code'] in inv:
            out.setdefault(inv[x['countryiso3code']], {})[int(x['date'])] = x['value']
    return out


def oms(code):
    d = json.loads(get("https://ghoapi.azureedge.net/api/%s?$filter=SpatialDim%%20in%%20(%s)" % (code, ','.join("'%s'" % ISO3[g] for g in UE))))
    inv = {v: k for k, v in ISO3.items()}
    out = {}
    for x in d['value']:
        # Une valeur 0 signifie « non déclaré » dans WUENIC (comme pour la France sur la page).
        if x['SpatialDim'] in inv and x['NumericValue']:
            out.setdefault(inv[x['SpatialDim']], {})[int(x['TimeDim'])] = x['NumericValue']
    return out


def ocde(flux, filtre):
    txt = get('https://sdmx.oecd.org/public/rest/data/%s/all?startPeriod=1995&format=csvfilewithlabels' % flux).decode('utf-8-sig')
    inv = {v: k for k, v in ISO3.items()}
    out = {}
    for x in csv.DictReader(io.StringIO(txt)):
        if x['REF_AREA'] in inv and x['OBS_VALUE'] and all(x.get(k) == v for k, v in filtre.items()):
            out.setdefault(inv[x['REF_AREA']], {})[int(x['TIME_PERIOD'])] = float(x['OBS_VALUE'])
    return out


def essais_cliniques(ans):
    """Même relevé que la France (calendrier san_essais) : essais interventionnels ayant un centre dans le pays, par année de début."""
    def n(g, y):
        u = ('https://clinicaltrials.gov/api/v2/studies?query.locn=%s&filter.advanced=AREA%%5BStudyType%%5DINTERVENTIONAL%%20AND%%20'
             'AREA%%5BStartDate%%5DRANGE%%5B%d-01-01,%d-12-31%%5D&countTotal=true&pageSize=1&fields=NCTId') % (NOM_EN[g], y, y)
        return g, y, json.loads(get(u))['totalCount']
    with ThreadPoolExecutor(4) as ex:
        res = list(ex.map(lambda a: n(*a), [(g, y) for g in UE for y in ans]))
    out = {}
    for g, y, v in res:
        out.setdefault(g, {})[y] = v
    return out


# Chaque indicateur : chemin dans data.json, périodes (une par table de la source), numérateur, dénominateur, facteur,
# et la phrase de méthode affichée sous la courbe. Les filtres sont ceux de la série France (voir donnees/calendrier.json).
def crim(iccs):
    def f():
        a = eurostat('crim_gen', 'unit=NR&iccs=' + iccs)
        b = eurostat('crim_off_cat', 'unit=NR&iccs=' + iccs)
        return {g: dict(list({y: v for y, v in a.get(g, {}).items() if y <= 2007}.items()) + list({y: v for y, v in b.get(g, {}).items() if y >= 2008}.items())) for g in UE}
    return f

_pisa = {}
def pisa(dom):
    """Scores moyens PISA des pays de l'UE, toutes vagues : OCDE, PISA 2025 Results (Volume I), annexe B1,
    tableaux I.B1.2a.36 (sciences), .37 (compréhension de l'écrit), .38 (mathématiques) (StatLink mrq53f)."""
    def f():
        if not _pisa:
            import openpyxl, tempfile
            noms = {'Czech Republic': 'CZ', 'Slovak Republic': 'SK'}
            noms.update({v: k for k, v in NOM_EN.items()})
            with tempfile.NamedTemporaryFile(suffix='.xlsx') as t:
                t.write(get('https://stat.link/mrq53f')); t.flush()
                wb = openpyxl.load_workbook(t.name, read_only=True, data_only=True)
                for tab, k in (('Table I.B1.2a.36', 's'), ('Table I.B1.2a.37', 'r'), ('Table I.B1.2a.38', 'm')):
                    rows = list(wb[tab].iter_rows(values_only=True))
                    cols = [(i, int(str(h).split()[1])) for i, h in enumerate(rows[7]) if h and str(h).startswith('PISA ') and '\n' not in str(h)]
                    _pisa[k] = {noms[n]: {y: r[i] for i, y in cols if isinstance(r[i], (int, float))} for r in rows[9:] for n in [str(r[0] or '').strip().rstrip('*').strip()] if n in noms}
        return _pisa[dom]
    return f


_prisons = {}
def prisons(ind):
    def f():
        if not _prisons:
            _prisons['ACT'] = eurostat('crim_pris_cap', 'unit=NR&indic_cr=PRIS_ACT_CAP')
            _prisons['OFF'] = eurostat('crim_pris_cap', 'unit=NR&indic_cr=PRIS_OFF_CAP')
        return _prisons[ind]
    return f

_otan = {}
OTAN_URL = None  # édition lue par outils/defense.py (C.otan.url) : la moyenne de l'UE se calcule sur la même édition que la France
def otan_xlsx():
    """Classeur Excel de l'OTAN qui accompagne « Defence Expenditure of NATO Countries » (devenu « Defence Investment of
    NATO Countries » en 2026) : celui de OTAN_URL s'il est connu, sinon celui de l'année, sinon celui de l'année
    précédente. Renvoie (classeur, adresse)."""
    if not _otan:
        import openpyxl, tempfile
        ed = 'https://www.nato.int/content/dam/nato/webready/documents/finance/def-exp-%d-en.xlsx'
        for u in ([OTAN_URL] if OTAN_URL else []) + [ed % a for a in (date.today().year, date.today().year - 1)]:
            try:
                b = get(u, essais=2)
            except IOError:
                continue
            if b[:2] != b'PK':
                continue
            t = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False); t.write(b); t.close()
            _otan['wb'], _otan['url'] = openpyxl.load_workbook(t.name, data_only=True), u
            break
        else:
            raise IOError('classeur de l\'OTAN introuvable (%s)' % (OTAN_URL or 'def-exp-%d-en.xlsx' % date.today().year))
    return _otan['wb'], _otan['url']


def otan_tab(feuille, bloc=None, tous=False):
    """{pays: {année: valeur}} d'un tableau du classeur de l'OTAN ; bloc : début du sous-titre du bloc voulu
    (« Constant » pour les prix constants), sinon le premier bloc. Les années estimées (« 2025e ») sont gardées :
    la page les affiche sans les compter (champ est). tous : garder aussi les pays hors UE (États-Unis)."""
    wb, _ = otan_xlsx()
    rows = list(wb[feuille].iter_rows(values_only=True))
    i = 0
    if bloc:
        i = next(j for j, r in enumerate(rows) if isinstance(r[0], str) and r[0].startswith(bloc))
    i = next(j for j in range(i, len(rows)) if rows[j][1] == 2014)
    ans = [int(str(v).rstrip('e')) if v is not None else None for v in rows[i][1:]]
    noms = {v: k for k, v in NOM_EN.items()}
    noms.update({'Slovak Republic': 'SK', 'Czech Republic': 'CZ', 'United States': 'US'})
    out = {}
    for r in rows[i + 1:]:
        if r[0] is None or not isinstance(r[0], str):
            break
        n = r[0].split(' (')[0].rstrip('*').strip()
        g = noms.get(n)
        if g and (g in UE or tous):
            out[g] = {y: float(v) for y, v in zip(ans, r[1:]) if y and isinstance(v, (int, float))}
    return out


def otan(feuille, bloc=None):
    return lambda: otan_tab(feuille, bloc)


def otan_part(feuille):
    """Montant d'une catégorie (équipement…) en dollars constants : part publiée × dépense de défense à prix constants."""
    def f():
        p, d = otan_tab(feuille), otan_tab('Table 2', 'Constant')
        return {g: {y: v * d[g][y] / 100 for y, v in p[g].items() if y in d.get(g, {})} for g in p}
    return f


def cofog(fonction, operation='TE'):
    return lambda: eurostat('gov_10a_exp', 'unit=MIO_EUR&sector=S13&na_item=%s&cofog99=%s' % (operation, fonction))


def silc(table, filtres):
    """Enquête européenne sur les revenus et les conditions de vie (Eurostat) : taux des 27 pays."""
    return lambda: eurostat(table, filtres)


def pnud_idh():
    """Indice de développement humain des 27 pays (PNUD, séries complètes du dernier rapport, lu par outils/vie.py)."""
    import vie
    r = vie.pnud_csv()[0]
    inv = {v: k for k, v in ISO3.items()}
    out = {}
    for x in r:
        if x['iso3'] in inv:
            out[inv[x['iso3']]] = {int(c[4:]): float(v) for c, v in x.items() if re.fullmatch(r'hdi_\d{4}', c) and v}
    return out


def pib_euros():
    return eurostat('nama_10_gdp', 'unit=CP_MEUR&na_item=B1GQ')


INDICATEURS = [
    dict(cle='securite.hom', periodes=[(1995, 2007), (2008, None)], num=crim('ICCS0101'), den=population, facteur=1e5,
         methode="homicides des pays de l'UE rapportés à leur population au 1er janvier", source="Eurostat, crim_gen jusqu'en 2007 et crim_off_cat depuis 2008 ; population : demo_gind"),
    dict(cle='securite.rob', periodes=[(1995, 2007), (2008, None)], num=crim('ICCS0401'), den=population, facteur=1e5,
         methode="vols avec violence des pays de l'UE rapportés à leur population au 1er janvier", source="Eurostat, crim_gen jusqu'en 2007 et crim_off_cat depuis 2008 ; population : demo_gind"),
    dict(cle='securite.bur', periodes=[(1995, 2007), (2008, None)], num=crim('ICCS05012'), den=population, facteur=1e5,
         methode="cambriolages de logement des pays de l'UE rapportés à leur population au 1er janvier", source="Eurostat, crim_gen jusqu'en 2007 et crim_off_cat depuis 2008 ; population : demo_gind"),
    dict(cle='securite.occ', periodes=[(2008, None)], num=prisons('ACT'), den=prisons('OFF'), facteur=100,
         methode="détenus des pays de l'UE divisés par leurs places de prison", source="Eurostat, crim_pris_cap"),
    dict(cle='securite.places', periodes=[(2008, None)], num=prisons('OFF'), den=population, facteur=1e5,
         methode="places de prison des pays de l'UE rapportées à leur population au 1er janvier", source="Eurostat, crim_pris_cap ; population : demo_gind"),
    dict(cle='sante.alc', periodes=[(2005, None)], taux=lambda: ocde('OECD.ELS.HD,DSD_HEALTH_LVNG@DF_HEALTH_LVNG_AC,', dict(MEASURE='AC', AGE='Y_GE15', SEX='_T', UNIT_MEASURE='L_PS')),
         methode="consommation des pays de l'UE publiée par l'OCDE, pondérée par leur population", source="OCDE, consommation d'alcool ; population : Eurostat, demo_gind"),
    dict(cle='sante.essais', periodes=[(2008, None)], num=lambda: essais_cliniques(range(2008, date.today().year)), den=pop_banque_mondiale, facteur=1e6,
         methode="essais lancés dans les pays de l'UE rapportés à leur population, un essai mené dans plusieurs pays comptant dans chacun, comme pour la France",
         source="ClinicalTrials.gov, relevé du %s ; population : Banque mondiale" % jour(date.today())),
    dict(cle='sante.vac.rougeole', periodes=[(2005, None)], taux=lambda: oms('MCV2'),
         methode="couverture des pays de l'UE pondérée par leur population", source="OMS et Unicef, WUENIC ; population : Eurostat, demo_gind"),
    dict(cle='sante.vac.dtp3', periodes=[(2005, None)], taux=lambda: oms('WHS4_100'),
         methode="couverture des pays de l'UE pondérée par leur population", source="OMS et Unicef, WUENIC ; population : Eurostat, demo_gind"),
    dict(cle='sante.vac.hpv', periodes=[(2010, None)], taux=lambda: oms('SDGHPVRECEIVED'),
         methode="couverture des pays de l'UE pondérée par leur population", source="OMS et Unicef, WUENIC ; population : Eurostat, demo_gind"),
    dict(cle='sante.dep.sein', periodes=[(2005, None)], taux=lambda: ocde('OECD.ELS.HD,DSD_HEALTH_PROC@DF_SCREEN,', dict(CANCER_SITE='CICDMNBR', AGE='Y50T69', SEX='F', DATA_SRCE='PRG')),
         methode="participation des pays de l'UE publiée par l'OCDE, pondérée par leur population", source="OCDE, données des programmes de dépistage ; population : Eurostat, demo_gind"),
    dict(cle='sante.dep.colon', periodes=[(2005, None)], taux=lambda: ocde('OECD.ELS.HD,DSD_HEALTH_PROC@DF_SCREEN,', dict(CANCER_SITE='CICDNCRA', AGE='Y50T74', SEX='_T', DATA_SRCE='PRG')),
         methode="participation des pays de l'UE publiée par l'OCDE, pondérée par leur population", source="OCDE, données des programmes de dépistage ; population : Eurostat, demo_gind"),
    dict(cle='sante.dep.col', periodes=[(2005, None)], taux=lambda: ocde('OECD.ELS.HD,DSD_HEALTH_PROC@DF_SCREEN,', dict(CANCER_SITE='CICDCRVX', AGE='Y20T69', SEX='F', DATA_SRCE='PRG')),
         methode="participation des pays de l'UE publiée par l'OCDE, pondérée par leur population", source="OCDE, données des programmes de dépistage ; population : Eurostat, demo_gind"),
    # Éducation : PISA, par vague (la note compare l'évolution d'une vague à la suivante).
    dict(cle='education.PI.m', periodes=[(2003, None)], taux=pisa('m'), pas='vagues',
         methode="scores des pays de l'UE pondérés par leur population", source="OCDE, PISA 2025, tableau I.B1.2a.38 ; population : Eurostat, demo_gind"),
    dict(cle='education.PI.r', periodes=[(2000, None)], taux=pisa('r'), pas='vagues',
         methode="scores des pays de l'UE pondérés par leur population", source="OCDE, PISA 2025, tableau I.B1.2a.37 ; population : Eurostat, demo_gind"),
    dict(cle='education.PI.s', periodes=[(2006, None)], taux=pisa('s'), pas='vagues',
         methode="scores des pays de l'UE pondérés par leur population", source="OCDE, PISA 2025, tableau I.B1.2a.36 ; population : Eurostat, demo_gind"),
    # Sécurité et justice : dépense publique pour la police, les tribunaux et les prisons. Eurostat ne publie la moyenne de
    # l'UE que depuis 2010 ; avant, Concret la calcule sur les mêmes tables (depuis 2010, le calcul redonne la sienne).
    dict(cle='securite.dpol', periodes=[(1995, None)], num=cofog('GF0301'), den=pib_euros, facteur=100,
         pourquoi="Eurostat ne publie la sienne que depuis 2010, et ce calcul la redonne exactement",
         methode="dépense publique pour la police des pays de l'UE divisée par leur PIB", source="Eurostat, gov_10a_exp (GF0301) et nama_10_gdp"),
    dict(cle='securite.dtrib', periodes=[(1995, None)], num=cofog('GF0303'), den=pib_euros, facteur=100,
         pourquoi="Eurostat ne publie la sienne que depuis 2010, et ce calcul la redonne exactement",
         methode="dépense publique pour les tribunaux des pays de l'UE divisée par leur PIB", source="Eurostat, gov_10a_exp (GF0303) et nama_10_gdp"),
    dict(cle='securite.dpris', periodes=[(1995, None)], num=cofog('GF0304'), den=pib_euros, facteur=100,
         pourquoi="Eurostat ne publie la sienne que depuis 2010, et ce calcul la redonne exactement",
         methode="dépense publique pour les prisons des pays de l'UE divisée par leur PIB", source="Eurostat, gov_10a_exp (GF0304) et nama_10_gdp"),
    # Défense. Dépense publique (Eurostat) : Eurostat ne publie la moyenne de l'UE que depuis 2010 ; avant, Concret la
    # calcule sur les mêmes tables (depuis 2010, le calcul redonne celle d'Eurostat). OTAN : les 23 pays de l'UE membres
    # de l'OTAN (Autriche, Irlande, Chypre et Malte n'en sont pas membres et ne publient pas ces chiffres).
    dict(cle='defense.dep', periodes=[(1995, None)], num=cofog('GF02'), den=pib_euros, facteur=100,
         pourquoi="Eurostat ne publie la sienne que depuis 2010, et ce calcul la redonne exactement",
         methode="dépense publique de défense des pays de l'UE divisée par leur PIB", source="Eurostat, gov_10a_exp (GF02) et nama_10_gdp"),
    dict(cle='defense.inv', periodes=[(1995, None)], num=cofog('GF02', 'P51G'), den=pib_euros, facteur=100,
         pourquoi="Eurostat ne publie la sienne que depuis 2010, et ce calcul la redonne exactement",
         methode="investissement public de défense des pays de l'UE divisé par leur PIB", source="Eurostat, gov_10a_exp (GF02, formation brute de capital) et nama_10_gdp"),
    dict(cle='defense.otan', periodes=[(2014, None)], num=otan('Table 2', 'Constant'), den=otan('Table 5', 'Constant'), facteur=100,
         methode="dépense de défense des pays de l'UE membres de l'OTAN divisée par leur PIB, à prix constants", source="OTAN, classeur des chiffres de défense des pays de l'OTAN, tableaux 2 et 5"),
    dict(cle='defense.equip', periodes=[(2014, None)], num=otan_part('Table 8a'), den=otan('Table 2', 'Constant'), facteur=100,
         methode="dépense d'équipement des pays de l'UE membres de l'OTAN divisée par leur dépense de défense, à prix constants", source="OTAN, classeur des chiffres de défense des pays de l'OTAN, tableaux 2 et 8a"),
    dict(cle='defense.AFF.pers', periodes=[(2014, None)], num=otan('Table 7'), den=population, facteur=1e6,
         methode="militaires des pays de l'UE membres de l'OTAN rapportés à leur population au 1er janvier", source="OTAN, classeur des chiffres de défense des pays de l'OTAN, tableau 7 ; population : Eurostat, demo_gind"),
    # Vie des Français : enquête européenne sur les revenus et les conditions de vie. Eurostat ne publie la moyenne de l'UE
    # que depuis 2010 ; Concret la calcule depuis 2005 sur les mêmes chiffres (depuis 2010, à 0,2 point près de celle d'Eurostat).
    dict(cle='vie.arr', periodes=[(2004, None)], taux=silc('ilc_mdes05', 'hhcomp=TOTAL&rskpovth=TOTAL&unit=PC'), pourquoi="Eurostat ne publie la sienne que depuis 2010 ; ce calcul la redonne à 0,2 point près",
         methode="part des personnes en retard de paiement dans les pays de l'UE, pondérée par leur population", source="Eurostat, ilc_mdes05 ; population : demo_gind"),
    dict(cle='vie.s80', periodes=[(2004, None)], taux=silc('ilc_di11', 'unit=RAT&age=TOTAL&sex=T'), pourquoi="Eurostat ne publie la sienne que depuis 2010 ; ce calcul la redonne à 0,02 près",
         methode="écart de revenu de chaque pays de l'UE, pondéré par sa population", source="Eurostat, ilc_di11 ; population : demo_gind"),
    dict(cle='vie.surch', periodes=[(2004, None)], taux=silc('ilc_lvho07a', 'unit=PC&age=TOTAL&sex=T&rskpovth=TOTAL'), pourquoi="Eurostat ne publie la sienne que depuis 2010 ; ce calcul la redonne à 0,2 point près",
         methode="part des personnes dont le logement coûte plus de 40 % du revenu dans les pays de l'UE, pondérée par leur population", source="Eurostat, ilc_lvho07a ; population : demo_gind"),
    dict(cle='vie.surp', periodes=[(2004, None)], taux=silc('ilc_lvho05a', 'unit=PC&age=TOTAL&sex=T&rskpovth=TOTAL'), pourquoi="Eurostat ne publie la sienne que depuis 2010 ; ce calcul la redonne à 0,2 point près",
         methode="part des personnes en logement surpeuplé dans les pays de l'UE, pondérée par leur population", source="Eurostat, ilc_lvho05a ; population : demo_gind"),
    dict(cle='vie.froid', periodes=[(2004, None)], taux=silc('ilc_mdes01', 'hhcomp=TOTAL&rskpovth=TOTAL&unit=PC'), pourquoi="Eurostat ne publie la sienne que depuis 2010 ; ce calcul la redonne à 0,2 point près",
         methode="part des personnes qui ne peuvent pas chauffer correctement leur logement dans les pays de l'UE, pondérée par leur population", source="Eurostat, ilc_mdes01 ; population : demo_gind"),
    dict(cle='vie.iw', periodes=[(2004, None)], taux=silc('ilc_iw01', 'wstatus=EMP&age=Y_GE18&sex=T&unit=PC'), pourquoi="Eurostat ne publie la sienne que depuis 2010 ; ce calcul la redonne à 0,2 point près",
         methode="part des travailleurs pauvres dans les pays de l'UE, pondérée par leur population", source="Eurostat, ilc_iw01 ; population : demo_gind"),
    dict(cle='vie.AFF.idh', periodes=[(1995, None)], taux=pnud_idh, dec=3,
         methode="indice de chaque pays de l'UE, pondéré par sa population", source="PNUD, séries complètes de l'indice de développement humain ; population : Eurostat, demo_gind"),
]


def critere(o, cle):
    """Série visée par cle : page.critère (dans C), page.critère.élément (dans sous), ou chemin direct (education.PI.m)."""
    p, k = cle.split('.', 1)
    c = o['series'][p]
    for part in k.split('.'):
        c = c['C'][part] if 'C' in c and part in c['C'] else c['sous'][part] if part in c.get('sous', {}) else c[part]
    return c


def pct(x):
    return ('%.1f' % x).replace('.', ',').replace(',0', '')


def calcule(ind, ans, dec):
    """Moyenne UE de chaque année (pays qui publient l'année) et son évolution (pays qui publient l'année et la précédente).
    La moyenne est arrondie comme la série France (dec décimales) ; l'évolution, qui sert à noter, garde 3 décimales."""
    P = population()
    if 'taux' in ind:
        t = ind['taux']()
        num = {g: {y: v * P[g][y] for y, v in t.get(g, {}).items() if y in P.get(g, {})} for g in UE}
        den = {g: {y: P[g][y] for y in t.get(g, {}) if y in P.get(g, {})} for g in UE}
        fac = 1
    else:
        num, den, fac = ind['num'](), ind['den'](), ind['facteur']
    a_la = lambda g, y: num.get(g, {}).get(y) is not None and den.get(g, {}).get(y)
    couv = lambda pays, y: 100 * sum(P[g].get(y, 0) for g in pays) / sum(P[g].get(y, 0) for g in UE)
    moy = lambda pays, y: fac * sum(num[g][y] for g in pays) / sum(den[g][y] for g in pays)
    niv, ev, info = {}, {}, []
    for a, b in ind['periodes']:
        b = b or ans[-1]
        nb, cv, sous = [], [], []
        for y in [y for y in ans if a <= y <= b]:
            i = ans.index(y)
            av = ans[i - 1] if i > 0 and ans[i - 1] >= a else None  # année (ou vague PISA) précédente, dans la même période
            pays = [g for g in UE if a_la(g, y)]
            if pays and couv(pays, y) >= SEUIL:
                v = moy(pays, y); niv[y] = round(v, dec)
            if av is None:
                continue
            deux = [g for g in UE if a_la(g, y) and a_la(g, av)]
            if deux and couv(deux, y) >= SEUIL:
                ev[y] = round(moy(deux, y) - moy(deux, av), 3)
                nb.append(len(deux)); cv.append(couv(deux, y))
            elif any(a_la(g, y) for g in UE):
                sous.append(y)
        faits = [y for y in niv if a <= y <= b]
        fin = max([y for y in ev if a <= y <= b] + faits + [a])
        # Grands pays (au moins 2 % de la population de l'UE) absents certaines années : la courbe de la moyenne
        # en est un peu déplacée ces années-là (pas la note, qui compare les mêmes pays d'une année à l'autre).
        trous = {}
        if faits:
            for g in UE:
                manque = [y for y in ans if min(faits) <= y <= max(faits) and y in niv and not a_la(g, y)]
                if manque and len(manque) < len(faits) and P[g].get(manque[0], 0) >= 0.02 * sum(P[x].get(manque[0], 0) for x in UE):
                    trous[g] = manque
        info.append(dict(de=a, a=fin, pays_min=min(nb) if nb else 0, pays_max=max(nb) if nb else 0, couv_min=round(min(cv), 1) if cv else None, sous_seuil=[y for y in sous if y <= fin], trous=trous))
    return niv, ev, info


def note(ind, info):
    ok = [i for i in info if i['couv_min'] is not None]
    tete = ("Pointillés : moyenne de l'UE calculée par Concret (%s ; %s) ; %s. " % (ind['methode'], ind['source'], ind['pourquoi'])) if ind.get('pourquoi') else \
        "Pointillés : moyenne de l'UE calculée par Concret, la source n'en publiant pas (%s ; %s). " % (ind['methode'], ind['source'])
    if not ok:
        return tete + "Les pays qui publient ce chiffre deux années de suite font moins de 70 % de la population de l'UE : son évolution ne peut pas être calculée, rien n'est donc noté."
    nb = lambda i: ('%d pays' % i['pays_min']) if i['pays_min'] == i['pays_max'] else ('%d à %d pays' % (i['pays_min'], i['pays_max']))
    per = ' ; '.join(('%d-%d : ' % (i['de'], i['a']) if len(ok) > 1 else '') + ('%s, toute la population de l\'UE' % nb(i) if i['couv_min'] >= 99.95 else '%s, au moins %s %% de la population de l\'UE' % (nb(i), pct(i['couv_min']))) for i in ok)
    def annees(l):
        r, d = [], l[0]
        for i in range(1, len(l) + 1):
            if i == len(l) or l[i] != l[i - 1] + 1:
                r.append(str(d) if d == l[i - 1] else ('%d et %d' if l[i - 1] == d + 1 else '%d à %d') % (d, l[i - 1])); d = l[i] if i < len(l) else None
        return ', '.join(r)
    tr = {}
    for i in ok:
        for g, l in i.get('trous', {}).items():
            tr[g] = sorted(set(tr.get(g, []) + l))
    fin = (" La courbe de la moyenne bouge aussi quand un grand pays manque (%s) ; la note, elle, compare toujours les mêmes pays d'une année à l'autre." % ' ; '.join('%s : %s' % (NOM_FR[g], annees(l)) for g, l in tr.items())) if tr else ''
    # Affiché sans note : pas de phrase sur ce qui est compté.
    if '.AFF.' in ind['cle'] and 'idh' in ind['cle']:
        return tete + "Moyenne calculée sur %s." % per
    # Années que le seuil de 70 % écarte : dites nommément (règle de CLAUDE.md).
    ss = sorted(set(y for i in ok for y in i.get('sous_seuil', [])))
    if ss and ind.get('pas') != 'vagues':
        fin = (" En %s : non compté, les pays de l'UE qui publient ce chiffre cette année-là et la précédente font moins de 70 %% de sa population." % annees(ss)) + fin
    if ind.get('pas') == 'vagues':
        return tete + "Chaque évolution porte sur les pays qui ont participé aux deux vagues (%s) ; sous 70 %% de la population de l'UE, la moyenne de l'OCDE est gardée pour cette vague.%s" % (per, fin.replace("d'une année à l'autre", "d'une vague à l'autre"))
    return tete + "Chaque évolution porte sur les pays qui publient les deux années (%s) ; sous 70 %% de la population de l'UE, l'année n'est pas comptée.%s" % (per, fin)


def main():
    a = sys.argv[1:]
    ecrire = '--ecrire' in a
    seul = a[a.index('--seul') + 1].split(',') if '--seul' in a else None
    o = json.load(open(os.path.join(ROOT, 'data.json'), encoding='utf-8'))
    global OTAN_URL
    OTAN_URL = o['series'].get('defense', {}).get('C', {}).get('otan', {}).get('url')
    # Chaque indicateur à part : une source injoignable (l'OTAN, par exemple) ne bloque pas les autres pages ;
    # ses moyennes gardent leur valeur et l'erreur est signalée (code de sortie 1).
    calc, echecs = [], []
    for ind in INDICATEURS:
        if (seul and ind['cle'] not in seul) or ind['cle'].split('.')[0] not in o['series']:
            continue
        try:
            c = critere(o, ind['cle'])
            dec = max([len(repr(float(v)).split('.')[1].rstrip('0')) for v in c['fr'] if v is not None] + [1])
            calc.append((ind, c, calcule(ind, c['y'], ind.get('dec', min(dec, 2)))))
        except Exception as e:
            echecs.append('%s (%s : %s)' % (ind['cle'], type(e).__name__, e))
    for e in echecs:
        print('Moyenne UE ✗ calcul impossible pour %s : valeurs gardées, rien écrit pour cet indicateur.' % e)
    for ind, c, (niv, ev, info) in calc:
        nv, ne = [niv.get(y) for y in c['y']], [ev.get(y) for y in c['y']]
        print('%s : %s' % (ind['cle'], ' ; '.join('%d-%d %s-%s pays, couverture min %s %%%s' % (i['de'], i['a'], i['pays_min'], i['pays_max'], i['couv_min'], (', sous le seuil : ' + ','.join(map(str, i['sous_seuil']))) if i['sous_seuil'] else '') for i in info)))
        print('    moyenne', ' '.join('%d:%s' % (y, v) for y, v in zip(c['y'], nv) if v is not None))
        av = c.get('ue_calc')
        if av is not None and av != nv:
            print('    changé par rapport à data.json :', ', '.join('%d %s→%s' % (y, p, q) for y, p, q in zip(c['y'], av, nv) if p != q))
        if ecrire:
            c['ue_calc'], c['ue_calc_ev'], c['ue_calc_info'], c['ue_calc_note'], c['ref'] = nv, ne, info, note(ind, info), 'ue_calc'
    if ecrire:
        open(os.path.join(ROOT, 'data.json'), 'w', encoding='utf-8').write(json.dumps(o, indent=0, ensure_ascii=False))
        if not seul and not echecs:
            # Les sources se mettent à jour à des dates différentes : on recalcule tout chaque mois.
            cal = json.load(open(os.path.join(ROOT, 'donnees/calendrier.json'), encoding='utf-8'))
            e = next((x for x in cal['indicateurs'] if x['id'] == 'ue_moyennes'), None)
            if e:
                auj = date.today()
                e['derniere_verification'] = auj.isoformat()
                e['prochaine_verification'] = date.fromordinal(auj.toordinal() + 30).isoformat()
                open(os.path.join(ROOT, 'donnees/calendrier.json'), 'w', encoding='utf-8').write(json.dumps(cal, indent=1, ensure_ascii=False) + '\n')
        print('Écrit dans data.json' + ('' if seul or echecs else ' ; prochaine vérification dans 30 jours') + '.')
    if echecs:
        sys.exit(1)


if __name__ == '__main__':
    main()
